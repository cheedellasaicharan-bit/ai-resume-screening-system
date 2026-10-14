"""Resume text extraction (PDF / DOCX / TXT) and simple name detection."""
import io
import re
import zipfile
from pathlib import Path

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_FILE_BYTES = 5 * 1024 * 1024  # 5 MB per file


class ParseError(Exception):
    """Raised with a human-readable message when a file cannot be used."""


def validate_file(filename: str, data: bytes) -> str:
    """Check extension, size and file signature. Returns the lower-case extension."""
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ParseError(f"Unsupported file type '{ext or 'none'}'. Use PDF, DOCX or TXT.")
    if len(data) == 0:
        raise ParseError("File is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise ParseError(f"File is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB.")
    if ext == ".pdf" and not data.lstrip()[:5] == b"%PDF-":
        raise ParseError("File is not a valid PDF (missing PDF header).")
    if ext == ".docx" and not zipfile.is_zipfile(io.BytesIO(data)):
        raise ParseError("File is not a valid DOCX (corrupted or wrong format).")
    return ext


def _read_pdf(data: bytes) -> str:
    import fitz  # PyMuPDF
    try:
        with fitz.open(stream=data, filetype="pdf") as doc:
            return "\n".join(page.get_text() for page in doc)
    except Exception as exc:
        raise ParseError(f"Could not read PDF: {exc}") from exc


def _read_docx(data: bytes) -> str:
    import docx  # python-docx
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ParseError(f"Could not read DOCX: {exc}") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _read_txt(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def extract_text(filename: str, data: bytes) -> str:
    """Return the plain text of an uploaded resume or raise ParseError."""
    ext = validate_file(filename, data)
    reader = {".pdf": _read_pdf, ".docx": _read_docx, ".txt": _read_txt}[ext]
    text = reader(data).replace("\x00", "").strip()
    if len(text) < 30:
        raise ParseError("No readable text found (the file may be a scanned image or nearly empty).")
    return text


def extract_name(text: str):
    """Guess the candidate name from the first short line. Returns None if unsure."""
    for line in text.splitlines()[:6]:
        line = line.strip()
        if not line or "fictional" in line.lower() or "sample" in line.lower():
            continue
        if re.search(r"[\d@:/]", line) or len(line) > 40:
            continue
        words = line.split()
        if 1 < len(words) <= 4 and all(w[0].isalpha() for w in words):
            return line.title() if line.isupper() else line
    return None
